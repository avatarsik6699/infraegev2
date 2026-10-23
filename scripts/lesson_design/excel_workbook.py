"""Build the Excel course's problem book (Задачник) from the site lessons (Change 197).

Usage:
    uv run --no-project --with openpyxl python -I scripts/lesson_design/excel_workbook.py
    uv run --no-project --with openpyxl --with formulas python -I \
        scripts/lesson_design/excel_workbook.py --check

Since Change 197 the site lessons are the source (the stand and its
``workbook.py`` were removed in T259; they stay in git history). The
book is built from the live sheets of the lesson media
(``apps/web/src/entities/course/content/excel-*.media.ts``), read through Node:

- «Начало» — what is where;
- one sheet per sheet block of a lesson, named «N.K Name» (lesson number of the
  course and the block's sheet name) in the order the lesson shows them; a
  block's other sheets («Товары», «Магазины») are shared by name, so formulas
  such as ``Товары!$A$2:$D$5`` work as in the lesson;
- «Задачи» — data of the bank tasks that attach the book, yellow cells for the
  reader's formulas; «Ответы» — the same tasks solved by formulas.

The build is reproducible (fixed save time), so the same data give the same
bytes. Without ``--check`` the book is stored in the practice bank by its
sha256 and every task that attached the previous book gets the new one (the
old blob is removed); bank.json is edited in place, its formatting kept.

``--check`` builds in a temporary folder, recomputes every formula with the
``formulas`` library — an engine independent of sheet_eval — skipping its known
differences (KNOWN), compares the answer cells with the bank answers and fails
when the bank still attaches another build (a lesson sheet changed). Exit code
1 on a mismatch.
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from sheet_eval import (  # noqa: E402
    EMPTY,
    FUNCTIONS,
    Book,
    Err,
    address,
    display,
    function_names,
    parse_address,
    shift,
    to_english,
    typed,
)

BANK = REPO / "content/practice-bank/bank.json"
BANK_FILES = REPO / "content/practice-bank/files"
FILENAME = "excel-zadachnik.xlsx"
YELLOW = "FFF2CC"
HEAD = "F3F3F1"
FORMATS = {"%": "0%", "0,0%": "0.0%", "0,00": "0.00", "дата": "dd.mm.yyyy"}
STAMP = "2026-10-21T00:00:00Z"  # the book's fixed date: same data → same bytes

# Bank tasks that attach the book: (lesson, task id, data rows from A1, answers).
# An answer is (range, formula in Russian locale typed into the range's first
# cell and copied over the range, a cell to check, the bank task whose answer
# that cell gives or None).
PRACTICE = [
    (
        "2.1",
        "excel-first-formulas-t08",
        [
            ["Товар", "Цена", "Количество", "Стоимость"],
            ["Блокнот", 60, 3],
            ["Линейка", 25, 4],
            ["Ластик", 10, 6],
            ["Итого"],
        ],
        [
            ("D2:D4", "=B2*C2", "D2", None),
            ("D5", "=СУММ(D2:D4)", "D5", "excel-first-formulas-t08"),
        ],
    ),
    (
        "3.1",
        "excel-conditions-t08",
        [["Число", "Признак"], [12], [7], [20], [3], [15], [], ["Итого"]],
        [
            ("B2:B6", "=ЕСЛИ(И(A2>=5;A2<=15);1;0)", "B2", None),
            ("B8", "=СУММ(B2:B6)", "B8", "excel-conditions-t08"),
        ],
    ),
    (
        "3.3",
        "excel-conditional-count-t06",
        [
            ["Ч1", "Ч2", "Ч3", "Ч4", "Ч5", "Ч6"]
            + ["Част. 1", "Част. 2", "Част. 3", "Част. 4", "Част. 5", "Част. 6"]
            + ["Одиночные", "Повторы"],
            [6, 6, 1, 2, 3, 9],
        ],
        [
            ("G2:L2", "=СЧЁТЕСЛИ($A2:$F2;A2)", "G2", None),
            ("M2", "=СУММЕСЛИ(G2:L2;1;A2:F2)", "M2", "excel-conditional-count-t06"),
            ("N2", '=СУММЕСЛИ(G2:L2;">1";A2:F2)', "N2", None),
        ],
    ),
    (
        "4.3",
        "excel-neighbour-cells-t02",
        [["Расход", "Накоплено"], [50], [30], [20], [100]],
        [
            ("B2", "=A2", "B2", None),
            ("B3:B5", "=B2+A3", "B4", "excel-neighbour-cells-t02"),
        ],
    ),
    (
        "5.1",
        "excel-from-task-to-table-t08",
        [
            ["Балл", "Признак"],
            [45],
            [80],
            [62],
            [80],
            [30],
            [95],
            [],
            ["Признаками", "СЧЁТЕСЛИ", "Среднее"],
        ],
        [
            ("B2:B7", "=ЕСЛИ(A2>=60;1;0)", "B3", None),
            ("A10", "=СУММ(B2:B7)", "A10", "excel-from-task-to-table-t08"),
            ("B10", '=СЧЁТЕСЛИ(A2:A7;">=60")', "B10", "excel-from-task-to-table-t08"),
            ("C10", '=СРЗНАЧЕСЛИ(A2:A7;">=60")', "C10", "excel-from-task-to-table-t15"),
        ],
    ),
]
# formula cells the formulas library computes differently from Excel; the
# reason is checked by hand in Microsoft's documentation (map.md, Change 176)
KNOWN = {
    "СЖПРОБЕЛЫ": "formulas keeps inner spaces; Excel's TRIM leaves one",
    "ЗНАЧЕН": "formulas parses text in English locale",
    "ПРОМЕЖУТОЧНЫЕ.ИТОГИ": "rows hidden by a filter are plain hidden rows in the file",
    "&": "formulas joins numbers with a decimal point, Excel in Russian locale with a comma",
    "ДАТА": "formulas returns a date object for some date cells",
    "ЕЧИСЛО": "formulas treats a text of digits as a number; Excel's ISNUMBER does not",
}

# Lessons of the Excel course in order with their sheet blocks, as the site
# shows them: a block's place is the first use of its media key in the lesson.
SITE_SHEETS_JS = r"""
const content = new URL("apps/web/src/entities/course/content/", `file://${process.argv[1]}/`);
const { excelCoursePublication } = await import(new URL("course-publication.mjs", content));
const { existsSync, readFileSync } = await import("node:fs");
const out = [];
excelCoursePublication.modules.forEach((module, m) => {
  module.lessonPlan.forEach((lesson, k) => out.push({ number: `${m + 1}.${k + 1}`, id: lesson.id }));
});
for (const lesson of out) {
  lesson.blocks = [];
  for (const part of ["", ".paid"]) {
    const media = new URL(`${lesson.id}${part}.media.ts`, content);
    const text = new URL(`${lesson.id}${part}.lesson.tsx`, content);
    if (!existsSync(media)) continue;
    const { lessonMedia } = await import(media);
    const source = readFileSync(text, "utf8");
    const keys = [...source.matchAll(/lessonMedia\.(\w+)/g)].map((m) => m[1]);
    for (const key of [...new Set(keys)]) {
      if (lessonMedia[key]?.kind === "sheet") lesson.blocks.push(lessonMedia[key].sheets);
    }
  }
}
process.stdout.write(JSON.stringify(out));
"""


def site_sheets() -> list[dict]:
    raw = subprocess.check_output(
        ["node", "--input-type=module", "-e", SITE_SHEETS_JS, str(REPO)], cwd=REPO
    )
    return json.loads(raw)


def sheet_title(base: str, used: set[str]) -> str:
    title = base[:31]
    k = 2
    while title in used:
        suffix = f" {k}"
        title = base[: 31 - len(suffix)] + suffix
        k += 1
    used.add(title)
    return title


def unknown_function(text: str) -> bool:
    """A formula that shows #ИМЯ? on purpose (lesson 5.1): stored as text."""
    return text.startswith("=") and any(n not in FUNCTIONS for n in function_names(text))


def cell_value(text: str):
    if text.startswith("="):
        return to_english(text)
    v = typed(text)
    if v is EMPTY:
        return None
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def put(ws, r: int, c: int, text: str) -> None:
    cell = ws.cell(r, c, text if unknown_function(text) else cell_value(text))
    if unknown_function(text):
        cell.data_type = "s"


def relocate(formula: str, drow: int) -> str:
    """Move a whole task block down: every row number moves, with $ or without."""
    return re.sub(
        r"(\$?[A-Z]{1,3}\$?)(\d+)(?![\wА-Яа-яЁё(])",
        lambda m: f"{m.group(1)}{int(m.group(2)) + drow}",
        formula,
    )


def build(path: Path):
    from openpyxl import Workbook  # pyright: ignore[reportMissingModuleSource]
    from openpyxl.styles import Font, PatternFill  # pyright: ignore[reportMissingModuleSource]

    head_fill = PatternFill(fill_type="solid", start_color=HEAD, end_color=HEAD)
    yellow = PatternFill(fill_type="solid", start_color=YELLOW, end_color=YELLOW)
    wb = Workbook()
    start = wb.active
    assert start is not None
    start.title = "Начало"
    used = {"Начало", "Задачи", "Ответы"}
    shared: dict[str, dict[str, str]] = {}
    reference = Book()  # every sheet of the file, for --check
    index = []
    for lesson in site_sheets():
        for block in lesson["blocks"]:
            for k, part in enumerate(block):
                if k > 0:
                    # a block's other sheets are shared by name across the book
                    cells = shared.setdefault(part["name"], {})
                    for addr, text in part["inputs"].items():
                        if cells.get(addr, text) != text:
                            raise SystemExit(
                                f"sheet {part['name']}: {addr} differs between lessons"
                            )
                        cells[addr] = text
                    continue
                title = sheet_title(f"{lesson['number']} {part['name']}", used)
                ws = wb.create_sheet(title)
                index.append((lesson["number"], title))
                ref_sheet = reference.sheet(title)
                for addr, text in part["inputs"].items():
                    r, c = parse_address(addr)
                    put(ws, r, c, text)
                    ref_sheet.inputs[(r, c)] = text
                    kind = part["formats"].get(addr)
                    if kind:
                        ws.cell(r, c).number_format = FORMATS[kind]
                for r in part["filtered"] + part["hidden"]:
                    ws.row_dimensions[r].hidden = True
                ref_sheet.filtered, ref_sheet.hidden = set(part["filtered"]), set(part["hidden"])
                c1, c2 = part["cols"]
                for c in range(c1, c2 + 1):
                    if part["rows"][0] == 1:
                        ws.cell(1, c).fill = head_fill
                        ws.cell(1, c).font = Font(bold=True)
                    ws.column_dimensions[ws.cell(1, c).column_letter].width = 14
    for name, cells in shared.items():
        ws = wb.create_sheet(name)
        ref_sheet = reference.sheet(name)
        for addr, text in cells.items():
            r, c = parse_address(addr)
            put(ws, r, c, text)
            ref_sheet.inputs[(r, c)] = text
        for c in range(1, max(parse_address(a)[1] for a in cells) + 1):
            ws.cell(1, c).fill = head_fill
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 14
    practice = wb.create_sheet("Задачи", 1)
    answers = wb.create_sheet("Ответы", 2)
    ref_answers = reference.sheet("Ответы")
    titles = {e["task"]["id"]: e["task"]["title"] for e in json.loads(BANK.read_text())["tasks"]}
    row = 1
    placed = []
    for lesson, task_id, data, cells in PRACTICE:
        heading = f"Урок {lesson}. {titles[task_id]}"
        for ws in (practice, answers):
            ws.cell(row, 1, heading).font = Font(bold=True)
        top = row + 1
        for i, values in enumerate(data):
            for j, v in enumerate(values):
                if v == "":
                    continue
                for ws in (practice, answers):
                    ws.cell(top + i, 1 + j, v)
                ref_answers.inputs[(top + i, 1 + j)] = str(v)
        for rng, formula, check_addr, answer_of in cells:
            first, _, last = rng.partition(":")
            r1, c1 = parse_address(first)
            r2, c2 = parse_address(last or first)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    copy = relocate(shift(formula, r - r1, c - c1), top - 1)
                    practice.cell(top + r - 1, c).fill = yellow
                    answers.cell(top + r - 1, c, to_english(copy))
                    ref_answers.inputs[(top + r - 1, c)] = copy
            cr, cc = parse_address(check_addr)
            placed.append((lesson, address(top + cr - 1, cc), answer_of))
        row = top + len(data) + 2
    for ws in (practice, answers):
        ws.column_dimensions["A"].width = 16
    start["A1"] = "Задачник к курсу «Excel с нуля для ЕГЭ»"
    start["A1"].font = Font(bold=True, size=14)
    start["A3"] = "Задачник собран из листов уроков: числа в уроках и в задачнике совпадают."
    start["A4"] = (
        "Листы «N.K …» — примеры урока N.K: меняй исходные числа и предсказывай результат."
    )
    start["A5"] = (
        "«Задачи» — данные задач уроков, жёлтые ячейки — место для твоих формул; «Ответы» — решения формулами."
    )
    start["A6"] = "Формулы записаны в файле по-английски, Excel покажет их на своём языке."
    start["A8"] = "Урок"
    start["B8"] = "Лист"
    for k, (lesson, title) in enumerate(index, start=9):
        start.cell(k, 1, lesson)
        start.cell(k, 2, title)
    start.column_dimensions["B"].width = 28
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    reproducible(path)
    return reference, placed


def reproducible(path: Path) -> None:
    """Drop the save time from the file: openpyxl stamps docProps/core.xml and
    every zip entry with «now». The bank stores the book by its sha256, so a
    rebuild of unchanged data must give the same bytes."""
    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    lambda m: m.group(1) + STAMP.encode(),
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)


def attached() -> set[str]:
    """Checksums of the books the bank tasks attach now."""
    bank = json.loads(BANK.read_text())
    return {
        f["checksum"]
        for e in bank["tasks"]
        for f in e["task"]["files"]
        if f["filename"] == FILENAME
    }


def bank_answers() -> dict[str, str]:
    bank = json.loads(BANK.read_text())
    return {e["task"]["id"]: e["task"]["checker"]["answer_variants"][0] for e in bank["tasks"]}


def check(path: Path, reference: Book, placed: list) -> list[str]:
    import formulas  # pyright: ignore[reportMissingImports]  # uv run --with formulas

    problems = []
    answers = bank_answers()
    for lesson, addr, answer_of in placed:
        if answer_of is None:
            continue
        got = display(reference.value("Ответы", addr))
        if got != answers[answer_of]:
            problems.append(f"Ответы!{addr} (урок {lesson}): {got}, bank {answers[answer_of]}")
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    if attached() != {sha}:
        problems.append(f"bank attaches {sorted(attached())}, the build is {sha}: rebuild the book")
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    compared = skipped = 0
    for name, sheet in reference.sheets.items():
        for (r, c), text in sheet.inputs.items():
            if not text.startswith("=") or unknown_function(text):
                continue
            if any(key in text for key in KNOWN):
                skipped += 1
                continue
            key = f"'[{path.name}]{name.upper()}'!{address(r, c)}"
            if key not in solution:
                problems.append(f"{name}!{address(r, c)}: not computed by formulas")
                continue
            other = solution[key].value
            other = other[0, 0] if hasattr(other, "shape") else other
            mine = reference.value(name, address(r, c))
            compared += 1
            if not agree(mine, other):
                problems.append(
                    f"{name}!{address(r, c)} {text}: sheet_eval {display(mine)}, formulas {other!r}"
                )
    print(f"formulas library: {compared} cells agree-checked, {skipped} skipped (KNOWN)")
    return problems


def agree(mine, other) -> bool:
    if isinstance(mine, Err):
        return type(other).__name__ == "XlError"
    if isinstance(mine, bool):
        return other is mine or (hasattr(other, "item") and other.item() is mine)
    if mine is EMPTY:
        return other in (0, "", None)
    if isinstance(mine, float):
        try:
            return abs(float(other) - mine) < 1e-9
        except (TypeError, ValueError):
            return False
    if type(other).__name__ == "XlError":
        return False
    return str(other) == mine or (mine == "" and other in (0, None))


def store(path: Path) -> str:
    """Put the book into the bank by sha256 and move every attachment to it."""
    data = path.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    old = attached()
    (BANK_FILES / sha).write_bytes(data)
    text = BANK.read_text()
    for previous in old - {sha}:
        text = text.replace(previous, sha)
        (BANK_FILES / previous).unlink(missing_ok=True)
    # the file record keeps its size next to the checksum
    text = re.sub(
        rf'("checksum": "{sha}",\s*"mime_type": "[^"]+",\s*"size_bytes": )\d+',
        lambda m: f"{m.group(1)}{len(data)}",
        text,
    )
    BANK.write_text(text)
    return sha


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / FILENAME
        reference, placed = build(path)
        if "--check" in sys.argv:
            problems = check(path, reference, placed)
            for p in problems:
                print(p)
            print(f"{len(placed)} answer cells; {len(problems)} problems")
            sys.exit(1 if problems else 0)
        sha = store(path)
    print(f"{FILENAME}: {len(reference.sheets)} sheets, {len(placed)} answer cells → files/{sha}")


if __name__ == "__main__":
    main()
