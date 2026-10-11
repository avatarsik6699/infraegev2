"""Build the Excel course's problem book (Задачник) from the stand's own sheets.

Usage:
    uv run --no-project --with openpyxl python -I content/lesson-design/course-excel/workbook.py
    uv run --no-project --with openpyxl --with formulas python -I \
        content/lesson-design/course-excel/workbook.py --check

Writes ``files/excel-zadachnik.xlsx`` (decision of 2026-10-11: the workbook comes
from the same data as the lessons, so the numbers in the text and in the file
always agree):

- «Начало» — what is where;
- one sheet per sheet block of the lessons, named «NN Name» (lesson number and
  the block's sheet name); a block's other sheets («Товары», «Магазины») are
  shared by name, so formulas such as ``Товары!$A$2:$D$5`` work as in the lesson;
- «Задачи» — data of lesson tasks with yellow cells for the reader's formulas,
  «Ответы» — the same tasks solved by formulas.

The build is reproducible: the save time is replaced by a fixed date, so the
same data always give the same bytes (the bank keeps the file by sha256).

Formulas are stored as an .xlsx stores them: English names, ``,``, ``.``
(``sheet_eval.to_english``); Excel shows them in its own language.

``--check`` builds the file in a temporary folder and recomputes it with the
``formulas`` library — an engine independent of sheet_eval — comparing every
formula cell with sheet_eval; known differences of that library (KNOWN below)
are skipped with a reason. The answers of «Ответы» are compared with the
lesson texts. Exit code 1 on a mismatch.
"""

from __future__ import annotations

import json
import re
import sys
import tempfile
from pathlib import Path

D = Path(__file__).resolve().parent
REPO = D.parents[2]
sys.path.insert(0, str(REPO / "scripts" / "lesson_design"))

from sheet_eval import (  # noqa: E402
    EMPTY,
    FUNCTIONS,
    Book,
    Err,
    address,
    display,
    function_names,
    parse_address,
    shift,
    to_english,
    typed,
)
from sheetblock import blocks_in  # noqa: E402

OUT = D / "files" / "excel-zadachnik.xlsx"
YELLOW = "FFF2CC"
HEAD = "F3F3F1"
FORMATS = {"%": "0%", "0,0%": "0.0%", "0,00": "0.00", "дата": "dd.mm.yyyy"}

# Lesson tasks with data: (lesson, title, data rows from A1, answers). An answer
# is (range, formula in Russian locale typed into the range's first cell and
# copied over the range, a cell to check, its expected display, a fragment of
# the lesson's answer that states it).
PRACTICE = [
    (
        "04",
        "Задача 8: стоимость покупок и итог",
        [
            ["Товар", "Цена", "Количество", "Стоимость"],
            ["Блокнот", 60, 3],
            ["Линейка", 25, 4],
            ["Ластик", 10, 6],
            ["Итого"],
        ],
        [
            ("D2:D4", "=B2*C2", "D2", "180", "8. 180, 100 и 60, итог 340"),
            ("D5", "=D2+D3+D4", "D5", "340", "итог 340"),
        ],
    ),
    (
        "05",
        "Задача 8: таблица сложения одной формулой",
        [["", 1, 2, 3, 4], [1], [2], [3], [4]],
        [("B2:E5", "=$A2+B$1", "E5", "8", "значение 8")],
    ),
    (
        "07",
        "Задача 8: признак «от 5 до 15 включительно»",
        [["Число", "Признак"], [12], [7], [20], [3], [15], [], ["Итого"]],
        [
            ("B2:B6", "=ЕСЛИ(И(A2>=5;A2<=15);1;0)", "B2", "1", "8. Признаки 1, 1, 0, 0, 1"),
            ("B8", "=СУММ(B2:B6)", "B8", "3", "итог 3"),
        ],
    ),
    (
        "09",
        "Задача 6: одиночные и повторяющиеся элементы строки",
        [
            [
                "Ч1",
                "Ч2",
                "Ч3",
                "Ч4",
                "Ч5",
                "Ч6",
                "Част. 1",
                "Част. 2",
                "Част. 3",
                "Част. 4",
                "Част. 5",
                "Част. 6",
                "Одиночные",
                "Повторы",
            ],
            [6, 6, 1, 2, 3, 9],
        ],
        [
            ("G2:L2", "=СЧЁТЕСЛИ($A2:$F2;A2)", "G2", "2", None),
            ("M2", "=СУММЕСЛИ(G2:L2;1;A2:F2)", "M2", "15", "Одиночные 15"),
            ("N2", '=СУММЕСЛИ(G2:L2;">1";A2:F2)', "N2", "12", "повторяющиеся 12"),
        ],
    ),
    (
        "13",
        "Задача 2: накопленный итог расходов",
        [["Расход", "Накоплено"], [50], [30], [20], [100]],
        [
            ("B2", "=A2", "B2", "50", None),
            ("B3:B5", "=B2+A3", "B5", "200", "2. 50, 80, 100 и 200."),
        ],
    ),
    (
        "14",
        "Задача 8: баллы не меньше 60 двумя способами",
        [
            ["Балл", "Признак"],
            [45],
            [80],
            [62],
            [80],
            [30],
            [95],
            [],
            ["Признаками", "СЧЁТЕСЛИ", "Среднее"],
        ],
        [
            ("B2:B7", "=ЕСЛИ(A2>=60;1;0)", "B3", "1", None),
            ("A10", "=СУММ(B2:B7)", "A10", "4", "8. 4 балла"),
            ("B10", '=СЧЁТЕСЛИ(A2:A7;">=60")', "B10", "4", None),
            ("C10", '=СРЗНАЧЕСЛИ(A2:A7;">=60")', "C10", "79,25", "среднее 79,25"),
        ],
    ),
]

# formula cells the formulas library computes differently from Excel; the
# reason is checked by hand in Microsoft's documentation (map.md, Change 176)
KNOWN = {
    "СЖПРОБЕЛЫ": "formulas keeps inner spaces; Excel's TRIM leaves one",
    "ЗНАЧЕН": "formulas parses text in English locale",
    "ПРОМЕЖУТОЧНЫЕ.ИТОГИ": "rows hidden by a filter are plain hidden rows in the file",
    "&": "formulas joins numbers with a decimal point, Excel in Russian locale with a comma",
    "ДАТА": "formulas returns a date object for some date cells",
    "ЕЧИСЛО": "formulas treats a text of digits as a number; Excel's ISNUMBER does not",
}


def lesson_sheets() -> list[tuple[str, list]]:
    course = json.loads((D / "course.json").read_text(encoding="utf-8"))
    out = []
    n = 0
    for module in course["modules"]:
        for lesson in module["lessons"]:
            n += 1
            text = (D / lesson["source"]).read_text(encoding="utf-8")
            out.append((f"{n:02d}", [b for _, b in blocks_in(text)]))
    return out


def sheet_title(base: str, used: set[str]) -> str:
    title = base[:31]
    k = 2
    while title in used:
        suffix = f" {k}"
        title = base[: 31 - len(suffix)] + suffix
        k += 1
    used.add(title)
    return title


def unknown_function(text: str) -> bool:
    """A formula that shows #ИМЯ? on purpose (lesson 14): stored as text."""
    return text.startswith("=") and any(n not in FUNCTIONS for n in function_names(text))


def put(ws, r: int, c: int, text: str) -> None:
    cell = ws.cell(r, c, text if unknown_function(text) else cell_value(text))
    if unknown_function(text):
        cell.data_type = "s"


def cell_value(text: str):
    if text.startswith("="):
        return to_english(text)
    v = typed(text)
    if v is EMPTY:
        return None
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def build(path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    head_fill = PatternFill(fill_type="solid", start_color=HEAD, end_color=HEAD)
    yellow = PatternFill(fill_type="solid", start_color=YELLOW, end_color=YELLOW)
    wb = Workbook()
    start = wb.active
    assert start is not None
    start.title = "Начало"
    used = {"Начало", "Задачи", "Ответы"}
    shared: dict[str, dict] = {}
    reference = Book()  # every sheet of the file, for --check
    index = []
    for lesson, blocks in lesson_sheets():
        for block in blocks:
            for k, part in enumerate(block.sheets):
                if k > 0:
                    # a block's other sheets are shared by name across the book
                    cells = shared.setdefault(part.name, {})
                    for cell, text in part.inputs.items():
                        if cells.get(cell, text) != text:
                            raise SystemExit(
                                f"sheet {part.name}: {address(*cell)} differs between lessons"
                            )
                        cells[cell] = text
                    continue
                title = sheet_title(f"{lesson} {part.name}", used)
                ws = wb.create_sheet(title)
                index.append((lesson, title))
                ref_sheet = reference.sheet(title)
                for (r, c), text in part.inputs.items():
                    put(ws, r, c, text)
                    ref_sheet.inputs[(r, c)] = text
                    kind = part.formats.get((r, c))
                    if kind:
                        ws.cell(r, c).number_format = FORMATS[kind]
                for r in part.filtered | part.hidden:
                    ws.row_dimensions[r].hidden = True
                ref_sheet.filtered, ref_sheet.hidden = set(part.filtered), set(part.hidden)
                if min(part.rows, default=2) == 1:
                    for c in range(min(part.cols), max(part.cols) + 1):
                        ws.cell(1, c).fill = head_fill
                        ws.cell(1, c).font = Font(bold=True)
                for c in range(min(part.cols), max(part.cols) + 1):
                    ws.column_dimensions[ws.cell(1, c).column_letter].width = 14
    for name, cells in shared.items():
        ws = wb.create_sheet(name)
        ref_sheet = reference.sheet(name)
        for (r, c), text in cells.items():
            put(ws, r, c, text)
            ref_sheet.inputs[(r, c)] = text
        for c in range(1, max(c for _, c in cells) + 1):
            ws.cell(1, c).fill = head_fill
            ws.column_dimensions[ws.cell(1, c).column_letter].width = 14

    practice = wb.create_sheet("Задачи", 1)
    answers = wb.create_sheet("Ответы", 2)
    ref_answers = reference.sheet("Ответы")
    row = 1
    placed = []
    for lesson, title, data, cells in PRACTICE:
        practice.cell(row, 1, f"Урок {int(lesson)}. {title}").font = Font(bold=True)
        answers.cell(row, 1, f"Урок {int(lesson)}. {title}").font = Font(bold=True)
        top = row + 1
        for i, values in enumerate(data):
            for j, v in enumerate(values):
                if v == "":
                    continue
                for ws in (practice, answers):
                    ws.cell(top + i, 1 + j, v)
                ref_answers.inputs[(top + i, 1 + j)] = str(v)
        for rng, formula, check_addr, want, fragment in cells:
            first, _, last = rng.partition(":")
            r1, c1 = parse_address(first)
            r2, c2 = parse_address(last or first)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    copy = relocate(shift(formula, r - r1, c - c1), top - 1)
                    practice.cell(top + r - 1, c).fill = yellow
                    answers.cell(top + r - 1, c, to_english(copy))
                    ref_answers.inputs[(top + r - 1, c)] = copy
            cr, cc = parse_address(check_addr)
            placed.append((lesson, address(top + cr - 1, cc), want, fragment))
        row = top + len(data) + 2
    for ws in (practice, answers):
        ws.column_dimensions["A"].width = 16

    start["A1"] = "Задачник к курсу «Excel с нуля для ЕГЭ»"
    start["A1"].font = Font(bold=True, size=14)
    start["A3"] = "Задачник собран из листов уроков: числа в уроках и в задачнике совпадают."
    start["A4"] = "Листы «NN …» — примеры урока NN: меняй исходные числа и предсказывай результат."
    start["A5"] = (
        "«Задачи» — данные задач уроков, жёлтые ячейки — место для твоих формул; «Ответы» — решения формулами."
    )
    start["A6"] = "Формулы записаны в файле по-английски, Excel покажет их на своём языке."
    start["A8"] = "Урок"
    start["B8"] = "Лист"
    for k, (lesson, title) in enumerate(index, start=9):
        start.cell(k, 1, int(lesson))
        start.cell(k, 2, title)
    start.column_dimensions["B"].width = 28
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    reproducible(path)
    return reference, placed


STAMP = "2026-10-11T00:00:00Z"  # the book's fixed date: same data → same bytes


def reproducible(path: Path) -> None:
    """Drop the save time from the file: openpyxl stamps docProps/core.xml and
    every zip entry with «now». The bank stores the book by its sha256, so a
    rebuild of unchanged data must give the same bytes."""
    import zipfile

    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    lambda m: m.group(1) + STAMP.encode(),
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)


def relocate(formula: str, drow: int) -> str:
    """Move a whole task block down: every row number moves, with $ or without."""
    return re.sub(
        r"(\$?[A-Z]{1,3}\$?)(\d+)(?![\wА-Яа-яЁё(])",
        lambda m: f"{m.group(1)}{int(m.group(2)) + drow}",
        formula,
    )


def check(path: Path, reference: Book, placed: list) -> list[str]:
    import formulas  # pyright: ignore[reportMissingImports]  # uv run --with formulas

    problems = []
    for lesson, addr, want, fragment in placed:
        got = display(reference.value("Ответы", addr))
        if got != want:
            problems.append(f"Ответы!{addr} (урок {lesson}): {got}, expected {want}")
        if fragment:
            (lesson_file,) = D.glob(f"{lesson}-*.md")
            if fragment not in lesson_file.read_text(encoding="utf-8"):
                problems.append(f"урок {lesson}: {fragment!r} not in the lesson")
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    book_name = path.name
    compared = skipped = 0
    for name, sheet in reference.sheets.items():
        for (r, c), text in sheet.inputs.items():
            if not text.startswith("=") or unknown_function(text):
                continue
            reason = next((why for key, why in KNOWN.items() if key in text), None)
            if reason:
                skipped += 1
                continue
            key = f"'[{book_name}]{name.upper()}'!{address(r, c)}"
            if key not in solution:
                problems.append(f"{name}!{address(r, c)}: not computed by formulas")
                continue
            other = solution[key].value
            other = other[0, 0] if hasattr(other, "shape") else other
            mine = reference.value(name, address(r, c))
            compared += 1
            if not agree(mine, other):
                problems.append(
                    f"{name}!{address(r, c)} {text}: sheet_eval {display(mine)}, formulas {other!r}"
                )
    print(f"formulas library: {compared} cells agree-checked, {skipped} skipped (KNOWN)")
    return problems


def agree(mine, other) -> bool:
    if isinstance(mine, Err):
        return type(other).__name__ == "XlError"
    if isinstance(mine, bool):
        return other is mine or (hasattr(other, "item") and other.item() is mine)
    if mine is EMPTY:
        return other in (0, "", None)
    if isinstance(mine, float):
        try:
            return abs(float(other) - mine) < 1e-9
        except (TypeError, ValueError):
            return False
    if type(other).__name__ == "XlError":
        return False
    return str(other) == mine or (mine == "" and other in (0, None))


def main() -> None:
    if "--check" in sys.argv:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / OUT.name
            reference, placed = build(path)
            problems = check(path, reference, placed)
        for p in problems:
            print(p)
        print(f"{len(placed)} answer cells; {len(problems)} problems")
        sys.exit(1 if problems else 0)
    reference, placed = build(OUT)
    print(f"{OUT.relative_to(REPO)}: {len(reference.sheets)} sheets, {len(placed)} answer cells")


if __name__ == "__main__":
    main()
