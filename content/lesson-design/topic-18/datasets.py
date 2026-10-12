"""Generate the workbooks of topic 18 and print the answers of their tasks.

Usage:
    python3 -I content/lesson-design/topic-18/datasets.py
        print the answer of every exam-format task
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-18/datasets.py --write
        write files/18-N.xlsx
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-18/datasets.py --check
        every book holds exactly the numbers, walls and dark cells of its field

Every exam-format task works on our own field, never on a file of the collection.
As on the exam, a book holds the field from cell A1 without headers; walls are
thick cell borders (Change 181, T2) and the field has a thick frame. A wall
between two cells is written on both of them (right of the left cell and left
of the right one, below the upper cell and above the lower one): exam files use
either side. Cells the robot may not enter hold −1 and a grey fill.

Walls are drawn as letters «Г» like on the exam: a ``corner`` closes a cell on
the right and below (a finish of a robot moving right and down), an ``entry``
closes it on the left and above — the cells behind such a corner are cut off
from the start (a rectangle of unreachable cells) or, for a field with several
starts, the corner is a start. Fields come from fixed seeds and never change;
``assert_shape`` checks that each field has the property its tasks rely on.

The answers here come from a forward table, as in the lesson; checks.py
recomputes them by recursion from the finish and by the stand's spreadsheet
evaluator on the whole field.
"""

from __future__ import annotations

import random
import re
import sys
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
FILES = HERE / "files"

Cell = tuple[int, int]
Moves = list[tuple[int, int]]

RD: Moves = [(0, 1), (1, 0)]  # right, down
LD: Moves = [(0, -1), (1, 0)]  # left, down
DIAG: Moves = [(0, 1), (1, 0), (1, 1)]  # right, down, right-down


@dataclass
class Field:
    a: list[list[int]]
    right: set[Cell] = field(default_factory=set)  # a wall to the right of the cell
    down: set[Cell] = field(default_factory=set)  # a wall below the cell
    dark: set[Cell] = field(default_factory=set)  # the robot may not enter

    @property
    def h(self) -> int:
        return len(self.a)

    @property
    def w(self) -> int:
        return len(self.a[0])

    def can(self, r: int, c: int, dr: int, dc: int) -> bool:
        """May the robot step from (r, c) by (dr, dc)? Walls, edges, dark cells."""
        nr, nc = r + dr, c + dc
        if not (0 <= nr < self.h and 0 <= nc < self.w) or (nr, nc) in self.dark:
            return False
        if dr == 0 and dc == 1:
            return (r, c) not in self.right
        if dr == 0 and dc == -1:
            return (r, nc) not in self.right
        if dr == 1 and dc == 0:
            return (r, c) not in self.down
        if dr == -1 and dc == 0:
            return (nr, c) not in self.down
        return True  # diagonal steps: the fields of such tasks have no walls


def cell_name(r: int, c: int) -> str:
    return "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[c] + str(r + 1)


# --- solving: the forward table of the lesson ----------------------------------------------

Gain = Callable[[Field, Cell, Cell], int]


def coin(f: Field, p: Cell, q: Cell) -> int:
    return f.a[q[0]][q[1]]


def best_table(
    f: Field,
    pick=max,
    moves: Moves = RD,
    starts: tuple[Cell, ...] = ((0, 0),),
    gain: Gain = coin,
    first: Callable[[Field, Cell], int] | None = None,
) -> list[list[int | None]]:
    """Best total from a start to every cell; None — no route («X» in the sheet).

    Cells are taken so that every predecessor comes first: rows in the direction
    of the vertical move, inside a row in the direction of the horizontal one
    (a diagonal move never goes against both)."""
    down = any(dr > 0 for dr, _ in moves)
    rightwards = any(dc > 0 for dr, dc in moves if dr == 0)
    rows = range(f.h) if down else range(f.h - 1, -1, -1)
    cols = range(f.w) if rightwards else range(f.w - 1, -1, -1)
    t: list[list[int | None]] = [[None] * f.w for _ in range(f.h)]
    for r, c in starts:
        t[r][c] = first(f, (r, c)) if first else f.a[r][c]
    for r in rows:
        for c in cols:
            if (r, c) in starts or (r, c) in f.dark:
                continue
            cands = []
            for dr, dc in moves:
                pr, pc = r - dr, c - dc
                if 0 <= pr < f.h and 0 <= pc < f.w:
                    prev = t[pr][pc]
                    if prev is not None and f.can(pr, pc, dr, dc):
                        cands.append(prev + gain(f, (pr, pc), (r, c)))
            if cands:
                t[r][c] = pick(cands)
    return t


def stops(f: Field, t: list[list[int | None]], moves: Moves = RD) -> list[Cell]:
    """Reachable cells the robot cannot leave: «угловые» or final cells."""
    return [
        (r, c)
        for r in range(f.h)
        for c in range(f.w)
        if t[r][c] is not None and not any(f.can(r, c, dr, dc) for dr, dc in moves)
    ]


def values(t, cells: list[Cell]) -> list[int]:
    return [t[r][c] for r, c in cells]


def unreachable(f: Field, t) -> int:
    return sum(1 for r in range(f.h) for c in range(f.w) if t[r][c] is None)


def naive_table(f: Field, pick=max, moves: Moves = RD) -> list[list[int]]:
    """The mistake the lesson warns about: a cell with no arrival counts as 0."""
    t = [[0] * f.w for _ in range(f.h)]
    t[0][0] = f.a[0][0]
    for r in range(f.h):
        for c in range(f.w):
            if (r, c) == (0, 0):
                continue
            cands = [
                t[r - dr][c - dc]
                for dr, dc in moves
                if 0 <= r - dr and 0 <= c - dc and f.can(r - dr, c - dc, dr, dc)
            ]
            t[r][c] = f.a[r][c] + (pick(cands) if cands else 0)
    return t


# --- fields -------------------------------------------------------------------------


def coins(rnd: random.Random, h: int, w: int, lo: int = 1, hi: int = 100) -> list[list[int]]:
    return [[rnd.randint(lo, hi) for _ in range(w)] for _ in range(h)]


def corner(f: Field, r: int, c: int, up: int, left: int) -> None:
    """A «Г» closing (r, c) on the right and below: `up` cells of the vertical
    wall upwards from r, `left` cells of the horizontal wall leftwards from c."""
    for k in range(up):
        if r - k >= 0 and c < f.w - 1:
            f.right.add((r - k, c))
    for k in range(left):
        if c - k >= 0 and r < f.h - 1:
            f.down.add((r, c - k))


def entry(f: Field, r: int, c: int, down: int, right: int) -> None:
    """A «Г» closing (r, c) on the left and above: the vertical wall goes down
    from r on the left of column c, the horizontal one right from c above row r."""
    for k in range(down):
        if r + k < f.h and c > 0:
            f.right.add((r + k, c - 1))
    for k in range(right):
        if c + k < f.w and r > 0:
            f.down.add((r - 1, c + k))


def mirror(f: Field) -> Field:
    """The same field mirrored left to right (walls too): a «move right» field
    becomes a «move left» one."""
    w = f.w
    return Field(
        [row[::-1] for row in f.a],
        {(r, w - 2 - c) for r, c in f.right},
        {(r, w - 1 - c) for r, c in f.down},
        {(r, w - 1 - c) for r, c in f.dark},
    )


def field_1() -> Field:
    """15 × 15 without walls (subtopic 2)."""
    return Field(coins(random.Random(1801), 15, 15))


def field_2() -> Field:
    """20 × 20, walls closing cells on the right and below; finish in the corner (3)."""
    rnd = random.Random(1802)
    f = Field(coins(rnd, 20, 20))
    for r, c, up, left in [
        (5, 6, 4, 3),
        (9, 14, 6, 5),
        (14, 3, 5, 3),
        (17, 11, 7, 6),
        (12, 17, 4, 2),
    ]:
        corner(f, r, c, up, left)
    return f


def field_3() -> Field:
    """20 × 20, finishes in «угловые» cells, two cut-off rectangles (4, 5, итог)."""
    rnd = random.Random(1803)
    f = Field(coins(rnd, 20, 20))
    for r, c, up, left in [(9, 6, 5, 3), (10, 15, 6, 4), (15, 8, 4, 5), (18, 17, 3, 3)]:
        corner(f, r, c, up, left)
    for r, c, d, rr in [(3, 10, 5, 3), (12, 2, 4, 2)]:
        entry(f, r, c, d, rr)
    return f


def field_4() -> Field:
    """20 × 20, «угловые» finishes; the start is inside the field (4)."""
    rnd = random.Random(1804)
    f = Field(coins(rnd, 20, 20))
    for r, c, up, left in [(8, 9, 5, 4), (13, 16, 6, 3), (17, 6, 4, 4), (4, 14, 3, 3)]:
        corner(f, r, c, up, left)
    entry(f, 11, 12, 3, 2)
    return f


def field_5() -> Field:
    """20 × 20 for a robot moving left and down from the top right cell (5):
    a right-and-down field mirrored, so its «угловые» cells are closed left and below."""
    rnd = random.Random(1805)
    g = Field(coins(rnd, 20, 20))
    for r, c, up, left in [(7, 5, 4, 4), (11, 13, 5, 3), (16, 9, 6, 5), (5, 16, 3, 2)]:
        corner(g, r, c, up, left)
    entry(g, 13, 3, 3, 3)
    return mirror(g)


def field_6() -> Field:
    """20 × 20 with several starts closed on the left and above; one finish (5)."""
    rnd = random.Random(1806)
    f = Field(coins(rnd, 20, 20))
    for r, c, d, rr in [(4, 6, 5, 4), (8, 13, 4, 5), (13, 4, 4, 6), (15, 15, 3, 3), (2, 16, 3, 2)]:
        entry(f, r, c, d, rr)
    corner(f, 10, 9, 3, 3)
    return f


def field_7() -> Field:
    """15 × 15 integers −50…50, moves right, down and right-down (6)."""
    return Field(coins(random.Random(1807), 15, 15, -50, 50))


def field_8() -> Field:
    """15 × 15 naturals with dark cells (−1) the robot may not enter (6)."""
    rnd = random.Random(1808)
    f = Field(coins(rnd, 15, 15, 1, 99))
    for r, c in rnd.sample([(r, c) for r in range(15) for c in range(15) if r + c > 0], 34):
        f.dark.add((r, c))
        f.a[r][c] = -1
    f.dark.discard((14, 14))
    f.a[14][14] = rnd.randint(1, 99)
    return f


def field_9() -> Field:
    """15 × 15 heights of a hilly ground: a step only if the heights differ by at
    most 50 (8). Each height is close to the mean of its upper and left neighbours."""
    rnd = random.Random(1809)
    a = [[0] * 15 for _ in range(15)]
    for r in range(15):
        for c in range(15):
            near = [a[r - 1][c]] * (r > 0) + [a[r][c - 1]] * (c > 0)
            base = sum(near) // len(near) if near else 200
            a[r][c] = min(400, max(1, base + rnd.randint(-60, 60)))
    return Field(a)


def field_10() -> Field:
    """15 × 15 naturals: parity rules for steps, diagonal always allowed (7)."""
    return Field(coins(random.Random(1810), 15, 15, 10, 99))


FIELDS: dict[str, Callable[[], Field]] = {
    "18-1": field_1,
    "18-2": field_2,
    "18-3": field_3,
    "18-4": field_4,
    "18-5": field_5,
    "18-6": field_6,
    "18-7": field_7,
    "18-8": field_8,
    "18-9": field_9,
    "18-10": field_10,
}

START_4 = (5, 3)  # field 18-4: position (4; 6) — column 4, row 6


# --- the questions of the tasks -----------------------------------------------------


def fixed(f: Field) -> tuple[int, int]:
    mx, mn = best_table(f, max), best_table(f, min)
    return mx[-1][-1], mn[-1][-1]  # type: ignore[return-value]


def corners(f: Field, moves: Moves = RD, start: Cell = (0, 0)) -> tuple[int, int]:
    mx = best_table(f, max, moves, (start,))
    mn = best_table(f, min, moves, (start,))
    ends = stops(f, mx, moves)
    return max(values(mx, ends)), min(values(mn, ends))


def runs_to_finishes(f: Field, k: int) -> tuple[int, int]:
    mx, mn = best_table(f, max), best_table(f, min)
    ends = stops(f, mx)
    return sum(sorted(values(mx, ends))[-k:]), sum(sorted(values(mn, ends))[:k])


def start_cells(f: Field) -> list[Cell]:
    """Cells closed on the left and above (by a wall or the edge)."""
    return [
        (r, c)
        for r in range(f.h)
        for c in range(f.w)
        if not f.can(r, c, 0, -1) and not f.can(r, c, -1, 0)
    ]


def to_finish(f: Field, pick) -> list[list[int | None]]:
    """Best total from every cell to the bottom right cell, moving right and down:
    the same table for the reversed moves, started at the finish."""
    return best_table(f, pick, [(0, -1), (-1, 0)], ((f.h - 1, f.w - 1),))


def runs_from_starts(f: Field, k: int) -> tuple[int, int]:
    mx, mn = to_finish(f, max), to_finish(f, min)
    st = [s for s in start_cells(f) if mx[s[0]][s[1]] is not None]
    return sum(sorted(values(mx, st))[-k:]), sum(sorted(values(mn, st))[:k])


def energy(f: Field) -> tuple[int, int, int]:
    """Field with dark cells; a step costs |difference|, the start costs its number.
    Minimal reserve to reach some final cell, any final cell, by any route."""

    def step(f: Field, p: Cell, q: Cell) -> int:
        return abs(f.a[q[0]][q[1]] - f.a[p[0]][p[1]])

    mn = best_table(f, min, gain=step)
    mx = best_table(f, max, gain=step)
    ends = stops(f, mn)
    return min(values(mn, ends)), max(values(mn, ends)), max(values(mx, ends))


def routes_and_cut_off(f: Field, limit: int = 60) -> tuple[int, int]:
    """Number of routes to the bottom right cell when a step needs |Δheight| ≤ limit,
    and the number of cells the robot cannot reach."""

    def ok(p: Cell, q: Cell) -> bool:
        return abs(f.a[p[0]][p[1]] - f.a[q[0]][q[1]]) <= limit

    n = [[0] * f.w for _ in range(f.h)]
    n[0][0] = 1
    for r in range(f.h):
        for c in range(f.w):
            if (r, c) == (0, 0):
                continue
            n[r][c] = (n[r][c - 1] if c and ok((r, c - 1), (r, c)) else 0) + (
                n[r - 1][c] if r and ok((r - 1, c), (r, c)) else 0
            )
    return n[-1][-1], sum(1 for row in n for x in row if x == 0)


def parity_steps(f: Field) -> tuple[int, int]:
    """Right only to a number of the same parity, down only to the other parity,
    diagonal always: the best sum at the bottom right cell, unreachable cells."""

    g = Field(f.a)
    t: list[list[int | None]] = [[None] * g.w for _ in range(g.h)]
    t[0][0] = g.a[0][0]
    for r in range(g.h):
        for c in range(g.w):
            if (r, c) == (0, 0):
                continue
            x = g.a[r][c]
            cands = []
            if c and t[r][c - 1] is not None and g.a[r][c - 1] % 2 == x % 2:
                cands.append(t[r][c - 1])
            if r and t[r - 1][c] is not None and g.a[r - 1][c] % 2 != x % 2:
                cands.append(t[r - 1][c])
            if r and c and t[r - 1][c - 1] is not None:
                cands.append(t[r - 1][c - 1])
            if cands:
                t[r][c] = x + max(c2 for c2 in cands if c2 is not None)
    return t[-1][-1], sum(1 for row in t for v in row if v is None)  # type: ignore[return-value]


def diag(f: Field) -> tuple[int, int]:
    mx, mn = best_table(f, max, DIAG), best_table(f, min, DIAG)
    return mx[-1][-1], mn[-1][-1]  # type: ignore[return-value]


TASKS: dict[str, tuple[str, Callable[[], object]]] = {
    "2.8": ("18-1", lambda: fixed(FIELDS["18-1"]())),
    "3.7": ("18-2", lambda: fixed(FIELDS["18-2"]())),
    "4.8": ("18-3", lambda: corners(FIELDS["18-3"]())),
    "4.9": ("18-4", lambda: corners(FIELDS["18-4"](), RD, START_4)),
    "5.7": ("18-5", lambda: corners(FIELDS["18-5"](), LD, (0, 19))),
    "5.8": ("18-6", lambda: runs_from_starts(FIELDS["18-6"](), 3)),
    "6.7": ("18-7", lambda: diag(FIELDS["18-7"]())),
    "6.8": ("18-8", lambda: energy(FIELDS["18-8"]())),
    "7.7": ("18-10", lambda: parity_steps(FIELDS["18-10"]())),
    "8.7": ("18-9", lambda: routes_and_cut_off(FIELDS["18-9"](), 50)),
    "F.1": ("18-3", lambda: runs_to_finishes(FIELDS["18-3"](), 2)),
    "F.2": ("18-2", lambda: corners(FIELDS["18-2"]())),
    "F.3": ("18-6", lambda: runs_from_starts(FIELDS["18-6"](), 2)),
    "F.4": ("18-1", lambda: diag(FIELDS["18-1"]())),
    "F.5": ("18-4", lambda: corners(FIELDS["18-4"]())),
}


def assert_shape() -> list[str]:
    """Each field has the property its tasks rely on."""
    problems = []
    f2 = FIELDS["18-2"]()
    if unreachable(f2, best_table(f2)) or fixed(f2) == fixed(Field(f2.a)):
        problems.append("18-2: walls must change the answer and leave every cell reachable")
    f3 = FIELDS["18-3"]()
    if not unreachable(f3, best_table(f3)):
        problems.append("18-3: no unreachable cells")
    naive = naive_table(f3, min)
    ends = stops(f3, best_table(f3))
    if min(naive[r][c] for r, c in ends) == corners(f3)[1]:
        problems.append("18-3: counting unreachable cells as 0 must spoil the minimum")
    f6 = FIELDS["18-6"]()
    if len(start_cells(f6)) < 5:
        problems.append("18-6: fewer than five starts")
    return problems


# --- books --------------------------------------------------------------------------------

STAMP = "2026-10-12T00:00:00Z"  # fixed save time: the same field gives the same bytes


def sides_of(f: Field) -> dict[Cell, set[str]]:
    """Walled sides of every cell: inner walls on both cells, the frame on the edge."""
    sides: dict[Cell, set[str]] = {}
    for r, c in f.right:
        sides.setdefault((r, c), set()).add("right")
        sides.setdefault((r, c + 1), set()).add("left")
    for r, c in f.down:
        sides.setdefault((r, c), set()).add("bottom")
        sides.setdefault((r + 1, c), set()).add("top")
    for r in range(f.h):
        sides.setdefault((r, 0), set()).add("left")
        sides.setdefault((r, f.w - 1), set()).add("right")
    for c in range(f.w):
        sides.setdefault((0, c), set()).add("top")
        sides.setdefault((f.h - 1, c), set()).add("bottom")
    return sides


def write_book(name: str, f: Field) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Border, PatternFill, Side

    thick = Side(style="thick", color="000000")
    grey = PatternFill("solid", start_color="808080", end_color="808080")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Лист1"
    for row in f.a:
        ws.append(row)
    for (r, c), sides in sides_of(f).items():
        ws.cell(r + 1, c + 1).border = Border(
            left=thick if "left" in sides else None,
            right=thick if "right" in sides else None,
            top=thick if "top" in sides else None,
            bottom=thick if "bottom" in sides else None,
        )
    for r, c in f.dark:
        ws.cell(r + 1, c + 1).fill = grey
    path = FILES / f"{name}.xlsx"
    wb.save(path)
    reproducible(path)


def reproducible(path: Path) -> None:
    """Drop the save time openpyxl stamps into the file (as course-excel/workbook.py)."""
    import zipfile

    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    lambda m: m.group(1) + STAMP.encode(),
                    data,
                )
            fixed_info = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed_info.compress_type = zipfile.ZIP_DEFLATED
            fixed_info.external_attr = info.external_attr
            out.writestr(fixed_info, data)


def read_book(name: str) -> Field:
    """The field as the book shows it: numbers, thick borders (either side), grey cells."""
    from openpyxl import load_workbook

    ws = load_workbook(FILES / f"{name}.xlsx").active
    assert ws is not None
    a = [
        [int(v) for v in row if isinstance(v, (int, float))]
        for row in ws.iter_rows(values_only=True)
    ]
    h, w = len(a), len(a[0])
    f = Field(a)

    def thick(r: int, c: int, side: str) -> bool:
        s = getattr(ws.cell(r + 1, c + 1).border, side)
        return s is not None and s.style == "thick"

    for r in range(h):
        for c in range(w):
            if c < w - 1 and (thick(r, c, "right") or thick(r, c + 1, "left")):
                f.right.add((r, c))
            if r < h - 1 and (thick(r, c, "bottom") or thick(r + 1, c, "top")):
                f.down.add((r, c))
            if ws.cell(r + 1, c + 1).fill.fill_type == "solid":
                f.dark.add((r, c))
    frame = all(thick(r, 0, "left") and thick(r, w - 1, "right") for r in range(h)) and all(
        thick(0, c, "top") and thick(h - 1, c, "bottom") for c in range(w)
    )
    if not frame:
        raise ValueError(f"{name}: the field has no frame")
    return f


def main() -> None:
    if "--write" in sys.argv:
        FILES.mkdir(exist_ok=True)
        for name, make in FIELDS.items():
            write_book(name, make())
        print(f"written: {len(FIELDS)} books in {FILES}")
    if "--check" in sys.argv:
        bad = []
        for name, make in FIELDS.items():
            f, g = make(), read_book(name)
            if (f.a, f.right, f.down, f.dark) != (g.a, g.right, g.down, g.dark):
                bad.append(name)
        print("books agree with the fields" if not bad else f"differ: {bad}")
        if bad:
            raise SystemExit(1)
    problems = assert_shape()
    for p in problems:
        print("SHAPE", p)
    for label, (name, solve) in TASKS.items():
        print(f"{label:5} {name:6} {solve()}")
    if problems:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
