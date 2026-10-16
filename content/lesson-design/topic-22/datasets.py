"""Generate the workbooks of topic 22 and print the answers of their tasks.

Usage:
    python3 -I content/lesson-design/topic-22/datasets.py
        print the answer of every exam-format task
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-22/datasets.py --write
        write files/22-N.xlsx
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-22/datasets.py --check
        every book holds exactly the table of its system

Every exam-format task works on our own system of processes, never on a file
of the collection. As on the exam, a book is one sheet with three columns from
A1: the process ID, its duration in ms and the IDs of the processes it depends
on — «0» when there are none, otherwise a list «3;7». Systems come from fixed
seeds and never change; ``assert_shape`` checks the property a task relies on.

The answers here come from the early schedule and the backward pass, as in the
lesson; checks.py recomputes them from the files with its own code.
"""

from __future__ import annotations

import random
import re
import sys
from dataclasses import dataclass
from pathlib import Path

HERE = Path(__file__).resolve().parent
FILES = HERE / "files"
STAMP = "2026-10-16T00:00:00Z"
HEAD = ["ID процесса B", "Время выполнения процесса B (мс)", "ID процесса(-ов) A"]


@dataclass
class System:
    rows: list[tuple[int, int, tuple[int, ...]]]

    def early(self) -> dict[int, tuple[int, int]]:
        out: dict[int, tuple[int, int]] = {}
        for i, d, deps in self.rows:  # suppliers always have smaller IDs
            s = max((out[j][1] for j in deps), default=0)
            out[i] = (s, s + d)
        return out

    def total(self) -> int:
        return max(f for _, f in self.early().values())

    def late_starts(self) -> dict[int, int]:
        """Latest start of every process keeping the minimal total time."""
        total = self.total()
        dur = {i: d for i, d, _ in self.rows}
        receivers: dict[int, list[int]] = {i: [] for i, _, _ in self.rows}
        for i, _, deps in self.rows:
            for j in deps:
                receivers[j].append(i)
        ls: dict[int, int] = {}
        for i, _, _ in reversed(self.rows):
            lf = min((ls[r] for r in receivers[i]), default=total)
            ls[i] = lf - dur[i]
        return ls


def make_system(seed: int, n: int, dmax: int) -> System:
    rng = random.Random(seed)
    rows = []
    for i in range(1, n + 1):
        k = 0 if i <= 3 else rng.choice([0, 1, 1, 2, 2, 3])
        deps = tuple(sorted(rng.sample(range(1, i), min(k, i - 1))))
        rows.append((i, rng.randint(1, dmax), deps))
    return System(rows)


BOOKS = {
    "22-1": (2201, 18, 20),
    "22-2": (2202, 20, 15),
    "22-3": (2203, 16, 12),
    "22-4": (2204, 20, 18),
    "22-5": (2205, 18, 16),
    "22-6": (2206, 30, 20),
}


def system(name: str) -> System:
    return make_system(*BOOKS[name])


def active(sched: dict[int, tuple[int, int]], t: int) -> int:
    return sum(1 for s, f in sched.values() if s <= t < f)


def task_1() -> int:
    """22-1: the minimal time to finish all processes."""
    return system("22-1").total()


MS_2 = 21


def task_2() -> int:
    """22-2: processes working during the 21st millisecond, i.e. on [20; 21)."""
    return active(system("22-2").early(), MS_2 - 1)


def task_3() -> tuple[int, int]:
    """22-3: the peak of simultaneous work and the longest segment with it."""
    sched = system("22-3").early()
    total = max(f for _, f in sched.values())
    counts = [active(sched, t) for t in range(total)]
    peak = max(counts)
    best = run = 0
    for c in counts:
        run = run + 1 if c == peak else 0
        best = max(best, run)
    return peak, best


K_4 = 10


def task_4() -> int:
    """22-4: the moment when at least 10 processes have finished."""
    return sorted(f for _, f in system("22-4").early().values())[K_4 - 1]


T_5 = 20


def task_5() -> int:
    """22-5: the largest number of processes that may start not earlier than
    20 ms while all processes finish in the minimal time."""
    return sum(1 for v in system("22-5").late_starts().values() if v >= T_5)


def task_6() -> int:
    """22-6: the minimal time to finish all of 30 processes."""
    return system("22-6").total()


ANSWERS = {
    "22-1": task_1,
    "22-2": task_2,
    "22-3": task_3,
    "22-4": task_4,
    "22-5": task_5,
    "22-6": task_6,
}


def assert_shape() -> None:
    p3, l3 = task_3()
    assert p3 >= 3 and l3 >= 2, (p3, l3)
    sched = system("22-4").early()
    fins = sorted(f for _, f in sched.values())
    assert fins[K_4 - 1] != fins[K_4 - 2], "the 10th finish must not coincide with the 9th"
    es = {i: s for i, (s, _) in system("22-5").early().items()}
    ls = system("22-5").late_starts()
    assert sum(1 for i in es if es[i] >= T_5) < sum(1 for i in ls if ls[i] >= T_5), (
        "late starts must matter"
    )


def deps_text(deps: tuple[int, ...]) -> int | str:
    return 0 if not deps else ";".join(map(str, deps))


def write_book(name: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Процессы"
    ws.append(HEAD)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, d, deps in system(name).rows:
        ws.append([i, d, deps_text(deps)])
    for letter, width in zip("ABC", (16, 34, 22), strict=True):
        ws.column_dimensions[letter].width = width
    path = FILES / f"{name}.xlsx"
    wb.save(path)
    reproducible(path)


def reproducible(path: Path) -> None:
    """Drop the save time openpyxl stamps into the file."""
    import zipfile

    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    lambda m: m.group(1) + STAMP.encode(),
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)


def check_book(name: str) -> bool:
    from openpyxl import load_workbook

    ws = load_workbook(FILES / f"{name}.xlsx").active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    want = [tuple(HEAD)] + [(i, d, deps_text(deps)) for i, d, deps in system(name).rows]
    return [tuple(r) for r in rows] == want


def main() -> None:
    assert_shape()
    if "--write" in sys.argv:
        FILES.mkdir(exist_ok=True)
        for name in BOOKS:
            write_book(name)
            print("written", name)
    elif "--check" in sys.argv:
        bad = [name for name in BOOKS if not check_book(name)]
        print("books ok" if not bad else f"books differ: {bad}")
        sys.exit(1 if bad else 0)
    else:
        for name, task in ANSWERS.items():
            print(name, task())


if __name__ == "__main__":
    main()
