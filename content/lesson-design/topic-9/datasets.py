"""Generate the data files of topic 9 and print the answers of their tasks.

Usage:
    python3 -I content/lesson-design/topic-9/datasets.py
        print the answer of every exam-format task (from the .txt files)
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-9/datasets.py --write
        write files/9-N.txt and files/9-N.xlsx (and the textbook file 9.txt)
    uv run --no-project --with openpyxl python -I content/lesson-design/topic-9/datasets.py --check
        the books hold exactly the numbers of the text files

Every exam-format task works on our own table, never on a file of the collection.
As on the exam, a task gives a workbook (numbers from cell A1, no header); the
same table saved as text (tab between numbers, one row per line) is given for a
program. Tables are generated from fixed seeds, so they never change.

Random rows almost never have the structure a task asks about («одно число
трижды, остальные различны»), so most tables are built row by row from a chosen
structure of repeats; some rows are planted on purpose: a comparison that is an
equality, a sorted row, a progression, a Pythagorean triple. checks.py recomputes
every answer by another method.
"""

from __future__ import annotations

import random
import re
import sys
from collections import Counter
from math import isqrt
from pathlib import Path

HERE = Path(__file__).resolve().parent
FILES = HERE / "files"
ROWS = 1000

# The textbook table of the lesson (source lines 1212–1219): six rows, answer 2.
TEXTBOOK = [
    [2, 2, 5, 7, 9, 11],
    [5, 5, 1, 3, 7, 9],
    [2, 2, 5, 5, 9, 11],
    [2, 2, 2, 7, 9, 11],
    [3, 3, 4, 5, 7, 8],
    [7, 7, 20, 22, 24, 26],
]

Row = list[int]


def structured(rnd: random.Random, sizes: tuple[int, ...], lo: int, hi: int) -> Row:
    """A row with groups of equal values of the given sizes, in random order."""
    values = rnd.sample(range(lo, hi + 1), len(sizes))
    row = [v for v, k in zip(values, sizes, strict=True) for _ in range(k)]
    rnd.shuffle(row)
    return row


def by_structure(seed: int, lo: int, hi: int, weights: dict[tuple[int, ...], int]) -> list[Row]:
    rnd = random.Random(seed)
    kinds, w = list(weights), list(weights.values())
    return [structured(rnd, rnd.choices(kinds, w)[0], lo, hi) for _ in range(ROWS)]


def uniform(seed: int, width: int, lo: int, hi: int) -> list[Row]:
    rnd = random.Random(seed)
    return [[rnd.randint(lo, hi) for _ in range(width)] for _ in range(ROWS)]


def plant(rows: list[Row], seed: int, count: int, make) -> list[Row]:
    """Replace `count` random rows by rows from make(rnd)."""
    rnd = random.Random(seed)
    for i in rnd.sample(range(len(rows)), count):
        rows[i] = make(rnd)
    return rows


def table_1() -> list[Row]:
    """Six naturals 1…99: sums, extremes, squares (subtopics 1, 2)."""
    return uniform(101, 6, 1, 99)


def table_2() -> list[Row]:
    """Five naturals 1…200, mostly different (subtopic 3)."""
    return by_structure(
        201, 1, 200, {(1, 1, 1, 1, 1): 75, (2, 1, 1, 1): 15, (2, 2, 1): 4, (3, 1, 1): 6}
    )


SIX = {
    (1, 1, 1, 1, 1, 1): 30,
    (2, 1, 1, 1, 1): 25,
    (2, 2, 1, 1): 15,
    (3, 1, 1, 1): 12,
    (2, 2, 2): 4,
    (3, 2, 1): 5,
    (3, 3): 2,
    (4, 1, 1): 4,
    (4, 2): 1,
    (5, 1): 1,
    (6,): 1,
}

SEVEN = {
    (1, 1, 1, 1, 1, 1, 1): 25,
    (2, 1, 1, 1, 1, 1): 20,
    (2, 2, 1, 1, 1): 15,
    (3, 1, 1, 1, 1): 12,
    (3, 2, 1, 1): 10,
    (4, 1, 1, 1): 6,
    (2, 2, 2, 1): 5,
    (3, 3, 1): 3,
    (3, 2, 2): 2,
    (4, 2, 1): 1,
    (5, 1, 1): 1,
}


def with_sum(rnd: random.Random, k: int, total: int, banned: set[int]) -> Row | None:
    """k different naturals 1…99 outside `banned` with the given sum, or None."""
    for _ in range(200):
        values = rnd.sample([v for v in range(1, 100) if v not in banned], k - 1)
        last = total - sum(values)
        if 1 <= last <= 99 and last not in values and last not in banned:
            return [*values, last]
    return None


def table_3() -> list[Row]:
    """Six naturals 1…99 with every structure of repeats (subtopics 3, 4).

    Planted on the border of the comparisons of the tasks: a triple r with
    singles summing to exactly 3r, a pair p with singles summing to exactly 4 · 2p.
    """

    def triple_equal(rnd: random.Random) -> Row:
        while True:
            r = rnd.randint(5, 60)
            singles = with_sum(rnd, 3, 3 * r, {r})
            if singles:
                row = [r, r, r, *singles]
                rnd.shuffle(row)
                return row

    def pair_equal(rnd: random.Random) -> Row:
        while True:
            p = rnd.randint(3, 24)
            singles = with_sum(rnd, 4, 8 * p, {p})
            if singles:
                row = [p, p, *singles]
                rnd.shuffle(row)
                return row

    rows = plant(by_structure(301, 1, 99, SIX), 302, 15, triple_equal)
    return plant(rows, 303, 15, pair_equal)


def table_4() -> list[Row]:
    """Seven naturals 1…99 with every structure of repeats (subtopics 3, 4)."""
    return by_structure(401, 1, 99, SEVEN)


def table_5() -> list[Row]:
    """Eight naturals 1…99: parity, numbers above the mean, quartets (5, 6)."""
    return uniform(501, 8, 1, 99)


def table_6() -> list[Row]:
    """Seven naturals 1…150; rows holding the integer part of their mean planted."""

    def with_mean(rnd: random.Random) -> Row:
        while True:
            row = [rnd.randint(1, 150) for _ in range(7)]
            row[rnd.randrange(7)] = sum(row) // 7
            if sum(row) // 7 in row:
                return row

    return plant(uniform(601, 7, 1, 150), 602, 60, with_mean)


def table_7() -> list[Row]:
    """Four naturals 1…60; progressions and equal pair sums planted (6, 7)."""

    def progression(rnd: random.Random) -> Row:
        d = rnd.randint(1, 15)
        a = rnd.randint(1, 60 - 3 * d)
        row = [a, a + d, a + 2 * d, a + 3 * d]
        rnd.shuffle(row)
        return row

    def pairs(rnd: random.Random) -> Row:
        while True:
            a, b, c = rnd.randint(1, 60), rnd.randint(1, 60), rnd.randint(1, 60)
            d = a + b - c
            if 1 <= d <= 60:
                row = [a, b, c, d]
                rnd.shuffle(row)
                return row

    rows = plant(uniform(701, 4, 1, 60), 702, 40, progression)
    return plant(rows, 703, 40, pairs)


def table_8() -> list[Row]:
    """Seven integers −50…50, structures of repeats; sorted rows planted (6)."""
    rows = by_structure(801, -50, 50, SEVEN)

    def ordered(rnd: random.Random) -> Row:
        kind = rnd.choice([(3, 1, 1, 1, 1), (2, 1, 1, 1, 1, 1), (1,) * 7, (2, 2, 1, 1, 1)])
        return sorted(structured(rnd, kind, -50, 50))

    return plant(rows, 802, 40, ordered)


def table_9() -> list[Row]:
    """Three naturals 1…100; Pythagorean triples planted (7)."""

    def right(rnd: random.Random) -> Row:
        p, q, r = rnd.choice([(3, 4, 5), (5, 12, 13), (8, 15, 17), (7, 24, 25), (20, 21, 29)])
        m = rnd.randint(1, 100 // r)
        row = [m * p, m * q, m * r]
        rnd.shuffle(row)
        return row

    return plant(uniform(901, 3, 1, 100), 902, 30, right)


def table_10() -> list[Row]:
    """Six naturals 1…40: frequencies by column and neighbours (8)."""
    return uniform(1001, 6, 1, 40)


TABLES = {
    "9-1": table_1,
    "9-2": table_2,
    "9-3": table_3,
    "9-4": table_4,
    "9-5": table_5,
    "9-6": table_6,
    "9-7": table_7,
    "9-8": table_8,
    "9-9": table_9,
    "9-10": table_10,
}


# --- the lesson's way of solving: one row at a time --------------------------------


def structure(a: Row) -> list[int]:
    return sorted(Counter(a).values())


def single(a: Row) -> Row:
    counts = Counter(a)
    return [x for x in a if counts[x] == 1]


def repeated(a: Row) -> Row:
    counts = Counter(a)
    return [x for x in a if counts[x] > 1]


def rest_sum(a: Row) -> int:
    return sum(a) - min(a) - max(a)


def is_square(x: int) -> bool:
    return x > 0 and isqrt(x) ** 2 == x


def count(rows: list[Row], check) -> int:
    return sum(1 for a in rows if check(a))


# subtopic 1 — sums and thresholds
def s1_7(rows):
    return count(rows, lambda a: sum(a) > 300 and max(a) < 90)


def s1_8(rows):
    return count(rows, lambda a: (max(a) > 95) != (min(a) < 5))


# subtopic 2 — sum, extremes, means
def s2_7(rows):
    return count(
        rows, lambda a: (min(a) + max(a)) ** 2 > sum(x * x for x in a) - min(a) ** 2 - max(a) ** 2
    )


def s2_8(rows):
    return count(rows, lambda a: 2 * (min(a) + max(a)) > rest_sum(a))


def s2_9(rows):
    return count(rows, lambda a: 5 * max(a) > 3 * (sum(a) - max(a)))


# subtopic 3 — repeats
def s3_7(rows):
    return count(rows, lambda a: len(set(a)) == len(a) and 2 * (min(a) + max(a)) <= rest_sum(a))


def s3_8(rows):
    return count(rows, lambda a: len(set(a)) < len(a) and a.count(max(a)) == 1)


def s3_9(rows):
    return count(rows, lambda a: structure(a) == [1, 1, 2, 3])


# subtopic 4 — groups
def s4_6(rows):
    def ok(a):
        r, s = repeated(a), single(a)
        return len(r) > 0 and len(s) > 0 and sum(r) * len(s) < sum(s) * len(r)

    return count(rows, ok)


def s4_7(rows):
    return count(
        rows, lambda a: structure(a) == [1, 1, 1, 3] and 3 * repeated(a)[0] >= sum(single(a))
    )


def s4_8(rows):
    return count(rows, lambda a: structure(a) == [1, 1, 2, 3] and min(repeated(a)) > min(single(a)))


def s4_9(rows):
    return count(
        rows, lambda a: structure(a) == [1, 1, 2, 2] and sum(set(repeated(a))) > sum(single(a))
    )


# subtopic 5 — parity and special numbers
def s5_7(rows):
    def ok(a):
        even = [x for x in a if x % 2 == 0]
        odd = [x for x in a if x % 2 != 0]
        return len(even) > len(odd) and sum(odd) > sum(even)

    return count(rows, ok)


def s5_8(rows):
    def ok(a):
        n, s = len(a), sum(a)
        above_even = [x for x in a if n * x > s and x % 2 == 0]
        above_odd = [x for x in a if n * x > s and x % 2 != 0]
        even = [x for x in a if x % 2 == 0]
        odd = [x for x in a if x % 2 != 0]
        return len(above_odd) > len(above_even) and sum(even) > sum(odd)

    return count(rows, ok)


def s5_9(rows):
    return count(rows, lambda a: sum(a) // len(a) in a and any(is_square(x) for x in a))


# subtopic 6 — order and the row number
def progression(a: Row) -> bool:
    b = sorted(a)
    return b[1] - b[0] == b[2] - b[1] == b[3] - b[2]


def s6_7(rows):
    return count(rows, lambda a: progression(a) and len(set(a)) == 4)


def s6_8(rows):
    def ok(a):
        b = sorted(a)
        return b[0] + b[3] == b[1] + b[2] and not progression(a)

    return count(rows, ok)


def s6_9(rows):
    answer = 0
    for number, a in enumerate(rows, start=1):
        if a == sorted(a) and structure(a) == [1, 1, 1, 1, 3]:
            answer = number
    return answer


def s6_10(rows):
    for a in rows:
        if structure(a) == [1, 1, 1, 1, 1, 2] and sum(single(a)) < repeated(a)[0]:
            return abs(min(a))
    return None


# subtopic 7 — triangles
def obtuse(a: Row) -> bool:
    b = sorted(a)
    return b[0] + b[1] > b[2] and b[0] ** 2 + b[1] ** 2 < b[2] ** 2


def s7_7(rows):
    return count(rows, lambda a: sorted(a)[0] ** 2 + sorted(a)[1] ** 2 == sorted(a)[2] ** 2)


def s7_8(rows):
    return count(rows, obtuse)


def s7_9(rows):
    return count(rows, lambda a: sorted(a)[0] + sorted(a)[1] < sorted(a)[3])


def s7_10(rows):
    def ok(a):
        p, q, r = sorted(a)
        return q * r > p * q + p * r

    return count(rows, ok)


# subtopic 8 — columns, the whole table, neighbours
def s8_7(rows):
    width = len(rows[0])
    columns = [Counter(a[j] for a in rows) for j in range(width)]
    answer = 0
    for a in rows:
        row_counts = Counter(a)
        good = sum(1 for j in range(width) if row_counts[a[j]] == 1 and columns[j][a[j]] - 1 >= 30)
        if good >= 2:
            answer += 1
    return answer


def s8_8(rows):
    every = Counter(x for a in rows for x in a)
    answer = 0
    for a in rows:
        row_counts = Counter(a)
        if any(row_counts[x] == 1 and every[x] - row_counts[x] == 160 for x in a):
            answer += 1
    return answer


def s8_9(rows):
    height, width = len(rows), len(rows[0])
    answer = 0
    for r, a in enumerate(rows):
        row_counts = Counter(a)
        good = 0
        for c in range(width):
            x = a[c]
            larger = False
            for dr in [-1, 0, 1]:
                for dc in [-1, 0, 1]:
                    nr, nc = r + dr, c + dc
                    if (dr, dc) != (0, 0) and 0 <= nr < height and 0 <= nc < width:
                        if rows[nr][nc] > x:
                            larger = True
            if row_counts[x] == 1 and not larger:
                good += 1
        if good >= 1:
            answer += 1
    return answer


# final page — mixed tasks
def f_1(rows):
    return count(
        rows, lambda a: structure(a) == [1, 1, 1, 1, 2] and sum(single(a)) <= 4 * sum(repeated(a))
    )


def f_2(rows):
    def ok(a):
        b = sorted(a)
        return b[-1] + b[-2] <= sum(b[:-2])

    return count(rows, ok)


def f_3(rows):
    answer = 0
    for number, a in enumerate(rows, start=1):
        if structure(a) == [1, 1, 1, 1, 3] and 4 * repeated(a)[0] > sum(single(a)):
            answer = number
    return answer


def f_4(rows):
    answer = 0
    for number, a in enumerate(rows, start=1):
        even = [x for x in a if x % 2 == 0]
        if len(set(a)) == len(a) and 2 * len(even) == len(a) and number % 2 == 0:
            answer += number
    return answer


def f_5(rows):
    answer = None
    for a in rows:
        if a == sorted(a) and len(set(a)) == len(a):
            answer = sum(a)
    return answer


TASKS = {
    "1.7": ("9-1", s1_7),
    "1.8": ("9-1", s1_8),
    "2.7": ("9-1", s2_7),
    "2.8": ("9-1", s2_8),
    "2.9": ("9-1", s2_9),
    "3.7": ("9-2", s3_7),
    "3.8": ("9-3", s3_8),
    "3.9": ("9-4", s3_9),
    "4.6": ("9", s4_6),
    "4.7": ("9-3", s4_7),
    "4.8": ("9-4", s4_8),
    "4.9": ("9-3", s4_9),
    "5.7": ("9-5", s5_7),
    "5.8": ("9-5", s5_8),
    "5.9": ("9-6", s5_9),
    "6.7": ("9-7", s6_7),
    "6.8": ("9-7", s6_8),
    "6.9": ("9-8", s6_9),
    "6.10": ("9-8", s6_10),
    "7.7": ("9-9", s7_7),
    "7.8": ("9-9", s7_8),
    "7.9": ("9-7", s7_9),
    "7.10": ("9-9", s7_10),
    "8.7": ("9-10", s8_7),
    "8.8": ("9-10", s8_8),
    "8.9": ("9-10", s8_9),
    "F.1": ("9-3", f_1),
    "F.2": ("9-1", f_2),
    "F.3": ("9-4", f_3),
    "F.4": ("9-5", f_4),
    "F.5": ("9-8", f_5),
}


# --- files -------------------------------------------------------------------------


def read(name: str) -> list[Row]:
    text = (FILES / f"{name}.txt").read_text(encoding="utf-8")
    return [[int(x) for x in line.split()] for line in text.splitlines() if line.strip()]


def write_text(name: str, rows: list[Row], sep: str = "\t") -> None:
    (FILES / f"{name}.txt").write_text(
        "".join(sep.join(map(str, a)) + "\n" for a in rows), encoding="utf-8"
    )


STAMP = "2026-10-12T00:00:00Z"  # fixed save time: the same table gives the same bytes


def write_book(name: str, rows: list[Row]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Лист1"
    for a in rows:
        ws.append(a)
    path = FILES / f"{name}.xlsx"
    wb.save(path)
    reproducible(path)


def reproducible(path: Path) -> None:
    """Drop the save time openpyxl stamps into the file (as course-excel/workbook.py)."""
    import zipfile

    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    lambda m: m.group(1) + STAMP.encode(),
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            out.writestr(fixed, data)


def book_rows(name: str) -> list[Row]:
    from openpyxl import load_workbook

    ws = load_workbook(FILES / f"{name}.xlsx", read_only=True).active
    assert ws is not None
    return [
        [int(v) for v in row if isinstance(v, (int, float))]
        for row in ws.iter_rows(values_only=True)
    ]


def main() -> None:
    if "--write" in sys.argv:
        FILES.mkdir(exist_ok=True)
        write_text("9", TEXTBOOK, " ")
        for name, make in TABLES.items():
            rows = make()
            write_text(name, rows)
            write_book(name, rows)
        print(f"written: 9.txt and {len(TABLES)} tables in {FILES}")
    if "--check" in sys.argv:
        bad = [n for n, make in TABLES.items() if not (book_rows(n) == read(n) == make())]
        print("books and texts agree" if not bad else f"differ: {bad}")
        if bad:
            raise SystemExit(1)
    for label, (name, solve) in TASKS.items():
        print(f"{label:5} {name:5} {solve(read(name))}")


if __name__ == "__main__":
    main()
